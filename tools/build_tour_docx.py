"""Собирает Word-версию экскурса по проекту из docs/onboarding/TOUR.md.

Исходник — Markdown: его читают на GitHub и правят вместе с кодом, а .docx пересобирается
этой командой. Разбирается только то подмножество Markdown, которым написан экскурс:
заголовки, абзацы, маркированные и нумерованные списки, таблицы, блоки кода, `код` и
**жирный** внутри строки. Встретится что-то другое — уйдёт в документ обычным текстом.

Запуск: python -X utf8 tools/build_tour_docx.py [исходник.md] [итог.docx]
Нужен python-docx (уже объявлен зависимостью сервера, на нём же работает docx_export).
"""
import pathlib
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = pathlib.Path(__file__).resolve().parent.parent
SRC = ROOT / "docs" / "onboarding" / "TOUR.md"
OUT = ROOT / "docs" / "onboarding" / "GradeBookAI-tour.docx"

BODY_FONT = "Calibri"
CODE_FONT = "Consolas"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)      # тёмно-синий, как у бренда Synapse
CODE_FILL = "F2F4F7"

INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")


def add_inline(par, text, size=None):
    """Строка с `кодом` и **жирным** — отдельными прогонами текста."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`") and len(part) > 1:
            run = par.add_run(part[1:-1])
            run.font.name = CODE_FONT
            run.font.size = Pt(9.5)
        elif part.startswith("**") and part.endswith("**") and len(part) > 3:
            run = par.add_run(part[2:-2])
            run.bold = True
        else:
            run = par.add_run(part)
        if size and not run.font.size:
            run.font.size = size


def shade(element, fill):
    """Заливка абзаца или ячейки. ShadingType.SOLID в Word даёт чёрный фон, поэтому clear."""
    props = element.get_or_add_pPr() if hasattr(element, "get_or_add_pPr") else element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def add_page_number_footer(section):
    par = section.footer.paragraphs[0]
    par.alignment = 1
    run = par.add_run("GradeBookAI изнутри · стр. ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    fld_run = par.add_run()
    fld_run.font.size = Pt(8)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        fld_run._r.append(el)


def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.2))
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        st = doc.styles[name]
        st.font.name = BODY_FONT
        st.font.size = Pt(size)
        st.font.color.rgb = ACCENT
        st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    add_page_number_footer(sec)
    props = doc.core_properties
    props.author = "Команда Synapse"
    props.last_modified_by = "Команда Synapse"
    props.title = "GradeBookAI изнутри: экскурс для программиста"
    props.comments = ""


def add_code(doc, lines):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(6)
    par.paragraph_format.line_spacing = 1.0
    par.paragraph_format.left_indent = Cm(0.3)
    shade(par._p, CODE_FILL)
    for i, line in enumerate(lines):
        run = par.add_run(line)
        run.font.name = CODE_FONT
        run.font.size = Pt(8.5)
        if i < len(lines) - 1:
            run.add_break()


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    width = max(len(r) for r in cells)
    table = doc.add_table(rows=len(cells), cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    # Ширина колонок — по самому длинному тексту в колонке, но не уже пятой части:
    # поровну делённая таблица отдаёт короткому имени файла полстраницы.
    longest = [max(len(r[j]) if j < len(r) else 0 for r in cells) for j in range(width)]
    shares = [max(n, sum(longest) / (width * 5)) for n in longest]
    total_cm = 16.6
    widths = [Cm(total_cm * s / sum(shares)) for s in shares]
    for i, row in enumerate(cells):
        for j in range(width):
            cell = table.cell(i, j)
            cell.width = widths[j]
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            add_inline(par, row[j] if j < len(row) else "", size=Pt(9.5))
            if i == 0:
                for run in par.runs:
                    run.bold = True
                shade(cell._tc, "E3E8EF")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_list_item(doc, text, marker, level=0):
    par = doc.add_paragraph()
    fmt = par.paragraph_format
    fmt.left_indent = Cm(0.9 + 0.6 * level)
    fmt.first_line_indent = Cm(-0.5)
    fmt.space_after = Pt(2)
    par.add_run(marker + "\t")
    fmt.tab_stops.add_tab_stop(Cm(0.9 + 0.6 * level))
    add_inline(par, text)


def build(src: pathlib.Path, out: pathlib.Path) -> None:
    lines = src.read_text(encoding="utf-8").replace("\r\n", "\n").split("\n")
    doc = Document()
    setup(doc)

    headings = [ln[3:].strip() for ln in lines if ln.startswith("## ")]
    i = 0
    para_buf = []
    toc_done = False

    def flush_para():
        if para_buf:
            add_inline(doc.add_paragraph(), " ".join(s.strip() for s in para_buf))
            para_buf.clear()

    while i < len(lines):
        ln = lines[i]
        if ln.startswith("```"):
            flush_para()
            j = i + 1
            while j < len(lines) and not lines[j].startswith("```"):
                j += 1
            add_code(doc, lines[i + 1:j])
            i = j + 1
            continue
        if ln.startswith("# "):
            flush_para()
            doc.add_paragraph(ln[2:].strip(), style="Title")
        elif ln.startswith("## "):
            flush_para()
            if not toc_done:
                # Содержание — перед первым разделом, по заголовкам исходника.
                doc.add_paragraph("Содержание", style="Heading 2")
                for h in headings:
                    p = doc.add_paragraph(h)
                    p.paragraph_format.left_indent = Cm(0.5)
                    p.paragraph_format.space_after = Pt(0)
                doc.add_page_break()
                toc_done = True
            doc.add_paragraph(ln[3:].strip(), style="Heading 1")
        elif ln.startswith("### "):
            flush_para()
            doc.add_paragraph(ln[4:].strip(), style="Heading 2")
        elif ln.startswith("|"):
            flush_para()
            j = i
            while j < len(lines) and lines[j].startswith("|"):
                j += 1
            add_table(doc, lines[i:j])
            i = j
            continue
        elif re.match(r"^\s*(- |\d+\. )", ln):
            flush_para()
            m = re.match(r"^(\s*)(- |(\d+)\. )(.*)$", ln)
            level = len(m.group(1)) // 2
            marker = "•" if m.group(2) == "- " else m.group(3) + "."
            text = [m.group(4)]
            # продолжение пункта — строки с отступом, пока не начался новый пункт
            while i + 1 < len(lines) and re.match(r"^ {2,}\S", lines[i + 1]) \
                    and not re.match(r"^\s*(- |\d+\. )", lines[i + 1]):
                i += 1
                text.append(lines[i].strip())
            add_list_item(doc, " ".join(text), marker, level)
        elif not ln.strip():
            flush_para()
        else:
            para_buf.append(ln)
        i += 1
    flush_para()
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))


if __name__ == "__main__":
    src = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else SRC
    out = pathlib.Path(sys.argv[2]) if len(sys.argv) > 2 else OUT
    build(src, out)
    print("собрано: %s (%d КБ)" % (out, out.stat().st_size // 1024))
