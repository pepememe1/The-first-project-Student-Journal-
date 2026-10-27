# -*- coding: utf-8 -*-
"""
test_account_ownership.py — учётка студента после выдачи пароля (27.10.2026, Влад).

━━ ПОРЯДОК, КОТОРЫЙ ЗДЕСЬ ДЕРЖИТСЯ ━━
Администратор выдаёт случайный пароль и вписывает почту студента. Студент входит, ему
предлагают привязать почту и сменить пароль; код уходит, только если почта совпала с
записью колледжа. После этого учётка ПРИНАДЛЕЖИТ СТУДЕНТУ: администратор не сбросит
пароль, не подменит почту колледжа и не возьмёт студента в лист выдачи. Забыл пароль —
код на почту, совпадающую И со своей подтверждённой, И с записью колледжа.

━━ ОБРАТНЫЙ ХОД ━━
Проверен откатом: снять вызов `owned_by_student` в `reset_credential` — краснеет
`test_onboarding_locks_the_account_from_the_admin`; убрать сверку с записью колледжа в
`password_recovery.start` — краснеет `test_recovery_code_needs_both_emails_to_match`;
`MAX_CODE_TRIES` не сжигает код — краснеет `test_code_burns_after_too_many_wrong_tries`.
"""
import re

import pytest

from conftest import make_admin

WEB = {"X-Client": "web", "User-Agent": "Mozilla/5.0 (Windows NT 10.0) Chrome/120.0"}
EMAIL = "ivanova@yandex.ru"
NEW_PW = "свойПароль42"


@pytest.fixture
def mail(monkeypatch):
    """Почта «настроена», письма складываются в ящик. Фоновую отправку кода делаем
    синхронной: иначе тест ждал бы поток и зависел бы от планировщика."""
    from app import mailer, password_recovery
    box = []
    monkeypatch.setattr(mailer, "configured", lambda: True)
    monkeypatch.setattr(mailer, "send_email",
                        lambda to, subject, body, html="": box.append(
                            {"to": to, "subject": subject, "body": body}) or True)

    class _Now:
        def __init__(self, target, daemon=True):
            self._t = target

        def start(self):
            self._t()

    monkeypatch.setattr(password_recovery.threading, "Thread", _Now)
    return box


@pytest.fixture
def data_key(monkeypatch):
    monkeypatch.setenv("GRADEBOOK_DATA_KEY", "11" * 32)


def _code(box, to=EMAIL):
    for letter in reversed(box):
        if letter["to"] == to:
            return re.search(r"Код: (\d{6})", letter["body"]).group(1)
    raise AssertionError(f"на {to} письма нет")


def _login(client, login, password):
    r = client.post("/auth/login", json={"login": login, "password": password}, headers=WEB)
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}", **WEB}


def _issued(client, admin, login="s1", group="К-1", admin_email=EMAIL):
    """Студент с выданным паролем и почтой, вписанной администратором."""
    r = client.post("/web/admin/students", json={
        "login": login, "surname": "Фам" + login, "name": "Имя", "group": group},
        headers=admin)
    assert r.status_code == 200, r.text
    r = client.post("/web/accounts/credential/reset", json={"login": login}, headers=admin)
    assert r.status_code == 200, r.text
    pw = r.json()["password"]
    if admin_email:
        r = client.post("/web/accounts/extra", json={"login": login, "admin_email": admin_email,
                                                     "admin_phone": ""}, headers=admin)
        assert r.status_code == 200, r.text
    return pw


def _onboard(client, me, box, email=EMAIL, new=NEW_PW):
    r = client.post("/me/onboarding/email", json={"email": email}, headers=me)
    assert r.status_code == 200, r.text
    r = client.post("/me/contacts/email/confirm", json={"code": _code(box, email)}, headers=me)
    assert r.status_code == 200, r.text
    r = client.post("/me/onboarding/password", json={"new": new, "confirm": new}, headers=me)
    assert r.status_code == 200, r.text


# ─────────────────────────────────────────────────────────────────────────────────────
# ПЕРВИЧНАЯ НАСТРОЙКА
# ─────────────────────────────────────────────────────────────────────────────────────

def test_onboarding_is_offered_while_the_issued_password_works(client, mail, data_key):
    admin = make_admin(client)
    pw = _issued(client, admin)
    me = _login(client, "s1", pw)
    view = client.get("/me/onboarding", headers=me).json()
    assert view["needed"] is True and view["step"] == "email" and view["admin_email_set"]


def test_code_goes_only_to_the_email_the_college_has(client, mail, data_key):
    admin = make_admin(client)
    me = _login(client, "s1", _issued(client, admin))
    r = client.post("/me/onboarding/email", json={"email": "thief@yandex.ru"}, headers=me)
    assert r.status_code == 400
    assert mail == [], "на чужую почту код уходить не должен"
    #Пароль без подтверждённой почты не меняется.
    r = client.post("/me/onboarding/password", json={"new": NEW_PW, "confirm": NEW_PW},
                    headers=me)
    assert r.status_code == 409


def test_without_college_email_onboarding_says_whom_to_ask(client, mail, data_key):
    admin = make_admin(client)
    me = _login(client, "s1", _issued(client, admin, admin_email=""))
    r = client.post("/me/onboarding/email", json={"email": EMAIL}, headers=me)
    assert r.status_code == 409 and r.headers.get("x-gb-reason") == "admin_email_missing"
    assert mail == []


def test_password_is_asked_twice_and_must_be_new(client, mail, data_key):
    admin = make_admin(client)
    pw = _issued(client, admin)
    me = _login(client, "s1", pw)
    client.post("/me/onboarding/email", json={"email": EMAIL}, headers=me)
    client.post("/me/contacts/email/confirm", json={"code": _code(mail)}, headers=me)
    r = client.post("/me/onboarding/password", json={"new": NEW_PW, "confirm": NEW_PW + "x"},
                    headers=me)
    assert r.status_code == 400 and "не совпадают" in r.json()["detail"]
    r = client.post("/me/onboarding/password", json={"new": pw, "confirm": pw}, headers=me)
    assert r.status_code == 400


def test_onboarding_locks_the_account_from_the_admin(client, mail, data_key):
    admin = make_admin(client)
    pw = _issued(client, admin)
    me = _login(client, "s1", pw)
    other = _login(client, "s1", pw)            #сессия того, кто держал лист группы
    _onboard(client, me, mail)

    assert client.get("/me/onboarding", headers=me).json()["needed"] is False
    #Сессия по выданному паролю закрыта, своя — жива.
    assert client.get("/me/contacts", headers=other).status_code == 401
    assert client.get("/me/contacts", headers=me).status_code == 200

    extra = client.get("/web/accounts/extra", params={"login": "s1"}, headers=admin).json()
    assert extra["owned"] is True and extra["credential"]["password"] == ""
    r = client.post("/web/accounts/credential/reset", json={"login": "s1"}, headers=admin)
    assert r.status_code == 409 and r.headers.get("x-gb-reason") == "account_owned"
    #Почту колледжа не подменить; телефон — можно.
    r = client.post("/web/accounts/extra", json={"login": "s1", "admin_email": "a@mail.ru",
                                                 "admin_phone": ""}, headers=admin)
    assert r.status_code == 409
    r = client.post("/web/accounts/extra", json={"login": "s1", "admin_email": EMAIL,
                                                 "admin_phone": "+7 900 111-22-33"},
                    headers=admin)
    assert r.status_code == 200, r.text
    #Пароль студента пережил попытку сброса.
    assert client.post("/auth/login", json={"login": "s1", "password": NEW_PW},
                       headers=WEB).status_code == 200
    assert client.post("/auth/login", json={"login": "s1", "password": pw},
                       headers=WEB).status_code == 401


def test_owned_student_is_never_in_the_rollout_sheet(client, mail, data_key):
    admin = make_admin(client)
    me = _login(client, "s1", _issued(client, admin))
    _issued(client, admin, login="s2", admin_email="")
    _onboard(client, me, mail)
    r = client.get("/web/accounts/rollout/students", params={"group": "К-1"}, headers=admin)
    states = {s["login"]: s["state"] for s in r.json()["students"]}
    assert states["s1"] == "owned" and states["s2"] == "issued"
    #Отметили только запертого — отказ, а не пустой лист.
    r = client.post("/web/accounts/rollout/export",
                    json={"group": "К-1", "student_ids": ["stud:s1"]}, headers=admin)
    assert r.status_code == 409
    #Галочку гасит страница, но запрос отправляется и руками: сервер пропускает сам.
    import io
    from openpyxl import load_workbook
    r = client.post("/web/accounts/rollout/export",
                    json={"group": "К-1", "student_ids": ["stud:s1", "stud:s2"],
                          "reset_existing": True}, headers=admin)
    assert r.status_code == 200, r.text
    ws = load_workbook(io.BytesIO(r.content)).active
    logins = {row[2] for row in ws.iter_rows(min_row=6, values_only=True)
              if row and isinstance(row[0], int)}
    assert logins == {"s2"}


def test_no_lock_while_email_is_not_the_colleges(client, mail, data_key):
    """Сменил пароль, но почта не подтверждена — восстановить нечем, замка нет."""
    admin = make_admin(client)
    pw = _issued(client, admin)
    me = _login(client, "s1", pw)
    r = client.post("/me/password", json={"current": pw, "new": NEW_PW}, headers=me)
    assert r.status_code == 200, r.text
    extra = client.get("/web/accounts/extra", params={"login": "s1"}, headers=admin).json()
    assert extra["owned"] is False
    assert client.post("/web/accounts/credential/reset", json={"login": "s1"},
                       headers=admin).status_code == 200


def test_onboarding_door_is_closed_once_the_issued_password_is_gone(client, mail, data_key):
    admin = make_admin(client)
    pw = _issued(client, admin)
    me = _login(client, "s1", pw)
    #Пароль, вписанный администратором, — тоже выданный (его знает колледж), поэтому
    #«свой» здесь получаем штатной сменой в настройках.
    assert client.post("/me/password", json={"current": pw, "new": NEW_PW},
                       headers=me).status_code == 200
    assert client.get("/me/onboarding", headers=me).json()["needed"] is False
    r = client.post("/me/onboarding/email", json={"email": EMAIL}, headers=me)
    assert r.status_code == 409


# ─────────────────────────────────────────────────────────────────────────────────────
# ВОССТАНОВЛЕНИЕ КОДОМ
# ─────────────────────────────────────────────────────────────────────────────────────

def _owned(client, admin, box):
    me = _login(client, "s1", _issued(client, admin))
    _onboard(client, me, box)
    box.clear()


def test_recovery_by_code_sets_a_new_password(client, mail, data_key):
    admin = make_admin(client)
    _owned(client, admin, mail)
    r = client.post("/auth/recover/code/start", json={"login": "s1", "email": EMAIL})
    assert r.status_code == 200 and r.json() == {"ok": True}
    code = _code(mail)
    bad = "000000" if code != "000000" else "111111"
    assert client.post("/auth/recover/code/verify",
                       json={"login": "s1", "code": bad}).status_code == 400
    r = client.post("/auth/recover/code/verify", json={"login": "s1", "code": code})
    assert r.status_code == 200, r.text
    token = r.json()["token"]
    #Код одноразовый.
    assert client.post("/auth/recover/code/verify",
                       json={"login": "s1", "code": code}).status_code == 400
    r = client.post("/auth/recover/confirm", json={"token": token, "password": "ещёНовый77"})
    assert r.status_code == 200, r.text
    assert client.post("/auth/login", json={"login": "s1", "password": NEW_PW},
                       headers=WEB).status_code == 401
    assert client.post("/auth/login", json={"login": "s1", "password": "ещёНовый77"},
                       headers=WEB).status_code == 200
    #Ключ тоже одноразовый.
    assert client.post("/auth/recover/confirm",
                       json={"token": token, "password": "третийПароль9"}).status_code == 400


def test_recovery_code_needs_both_emails_to_match(client, mail, data_key):
    admin = make_admin(client)
    #Почта колледжа есть, своя не подтверждена.
    _issued(client, admin, login="a1")
    #Своя подтверждена, но колледж записал другую.
    pw = _issued(client, admin, login="b1", admin_email="")
    me = _login(client, "b1", pw)
    client.post("/me/contacts/email/start", json={"email": EMAIL, "current": pw}, headers=me)
    client.post("/me/contacts/email/confirm", json={"code": _code(mail)}, headers=me)
    client.post("/web/accounts/extra", json={"login": "b1", "admin_email": "other@mail.ru",
                                             "admin_phone": ""}, headers=admin)
    mail.clear()
    answers = [client.post("/auth/recover/code/start", json={"login": login, "email": EMAIL})
               for login in ("a1", "b1", "nobody")]
    assert all(a.status_code == 200 and a.json() == {"ok": True} for a in answers)
    assert mail == [], "код ушёл на почту, не совпадающую с обеими записями"


def test_wrong_email_gets_the_same_answer_and_no_letter(client, mail, data_key):
    admin = make_admin(client)
    _owned(client, admin, mail)
    r = client.post("/auth/recover/code/start", json={"login": "s1", "email": "x@yandex.ru"})
    assert r.status_code == 200 and r.json() == {"ok": True}
    assert mail == []


def test_code_burns_after_too_many_wrong_tries(client, mail, data_key):
    from app import password_recovery, throttle
    admin = make_admin(client)
    _owned(client, admin, mail)
    client.post("/auth/recover/code/start", json={"login": "s1", "email": EMAIL})
    code = _code(mail)
    bad = "000000" if code != "000000" else "111111"
    for _ in range(password_recovery.MAX_CODE_TRIES):
        client.post("/auth/recover/code/verify", json={"login": "s1", "code": bad})
    throttle.reset()        #общий замок попыток — отдельная защита, здесь проверяем свою
    r = client.post("/auth/recover/code/verify", json={"login": "s1", "code": code})
    assert r.status_code == 400, "код обязан сгореть после серии неверных вводов"


def test_letters_are_not_sent_more_than_once_a_minute(client, mail, data_key):
    admin = make_admin(client)
    _owned(client, admin, mail)
    for _ in range(3):
        client.post("/auth/recover/code/start", json={"login": "s1", "email": EMAIL})
    assert len(mail) == 1


def test_student_editor_cannot_replace_an_owned_password(client, mail, data_key):
    """Поле «Новый пароль» в редакторе студента — тоже дверь, и она обходила весь замок:
    пароль снова «выданный», замок снят, почту колледжа можно подменить (Полковник,
    28.10.2026). ФИО при этом править можно — на нём вход не держится."""
    admin = make_admin(client)
    me = _login(client, "s1", _issued(client, admin))
    _onboard(client, me, mail)
    r = client.put("/web/admin/students/s1", json={"password": "adminPass99"}, headers=admin)
    assert r.status_code == 409 and r.headers.get("x-gb-reason") == "account_owned"
    r = client.put("/web/admin/students/s1", json={"surname": "Новая"}, headers=admin)
    assert r.status_code == 200, r.text
    assert client.post("/auth/login", json={"login": "s1", "password": "adminPass99"},
                       headers=WEB).status_code == 401
    extra = client.get("/web/accounts/extra", params={"login": "s1"}, headers=admin).json()
    assert extra["owned"] is True, "замок обязан пережить попытку"


def test_someone_elses_wrong_tries_do_not_burn_the_owners_code(client, mail, data_key):
    """Посторонний, знающий логин, не сжигает код владельцу: код сгорает для адреса
    после пяти его ошибок, для всех — после двадцати (Полковник, 28.10.2026)."""
    from app import password_recovery
    admin = make_admin(client)
    _owned(client, admin, mail)
    client.post("/auth/recover/code/start", json={"login": "s1", "email": EMAIL})
    code = _code(mail)
    bad = "000000" if code != "000000" else "111111"
    for _ in range(password_recovery.MAX_CODE_TRIES):
        client.post("/auth/recover/code/verify", json={"login": "s1", "code": bad},
                    headers={"X-Real-IP": "203.0.113.7"})
    r = client.post("/auth/recover/code/verify", json={"login": "s1", "code": code},
                    headers={"X-Real-IP": "203.0.113.7"})
    assert r.status_code == 400, "для перебиравшего адреса код сгорел"
    r = client.post("/auth/recover/code/verify", json={"login": "s1", "code": code},
                    headers={"X-Real-IP": "198.51.100.20"})
    assert r.status_code == 200 and r.json().get("token"), r.text
