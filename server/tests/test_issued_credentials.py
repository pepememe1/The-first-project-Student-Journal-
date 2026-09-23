# -*- coding: utf-8 -*-
"""test_issued_credentials.py — «Выкатить данные групп» и стартовые пароли (4.0).

Что здесь защищается:
  • РОЛЬ-СКОУП НА СЕРВЕРЕ. Куратор выдаёт данные только СВОИМ группам; подставленный
    чужой id студента молча отбрасывается, а не отдаёт чужой пароль.
  • ГРАНИЦА «ПРИДУМАННЫЙ САМИМ ПАРОЛЬ НЕ ВИДЕН НИКОМУ». Сменил пароль — стартовый
    стирается. Проверяется и штатный путь (/me/password), и путь, о котором модуль
    выдачи НЕ знает: сверка отпечатка хеша ловит любой.
  • ПОВТОРНАЯ ВЫГРУЗКА ДАЁТ ТЕ ЖЕ ПАРОЛИ — иначе вчерашняя распечатка врёт.
  • ВЫГРУЗКА НЕ ВЫБИВАЕТ ТЕХ, КТО УЖЕ ПОЛЬЗУЕТСЯ ЖУРНАЛОМ (пароль задан до 4.0).
  • ШИФРОТЕКСТ, А НЕ ПАРОЛЬ, В БАЗЕ; в журнале аудита пароля нет.
"""
import io

import pytest
from openpyxl import load_workbook

from conftest import make_admin

_DATA_KEY = "11" * 32


@pytest.fixture
def data_key(monkeypatch):
    """Ключ «Кузнечика» на время теста: без него `gost.encrypt` честно отдаёт текст как
    есть, и проверить «в базе шифротекст» было бы нечем."""
    monkeypatch.setenv("GRADEBOOK_DATA_KEY", _DATA_KEY)


def _student(client, admin, login, group, password=""):
    body = {"login": login, "surname": "Фам" + login, "name": "Имя", "group": group}
    if password:
        body["password"] = password
    r = client.post("/web/admin/students", json=body, headers=admin)
    assert r.status_code == 200, r.text


def _curator(client, admin, login, groups):
    r = client.post("/web/admin/teachers", json={
        "full_name": "Кур Атор", "login": login, "password": "curpass123",
        "subjects": [], "curated_groups": list(groups)}, headers=admin)
    assert r.status_code == 200, r.text
    r = client.post("/auth/login", json={"login": login, "password": "curpass123"})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}", "X-Client": "web"}


def _ids(client, headers, group):
    r = client.get("/web/accounts/rollout/students", params={"group": group}, headers=headers)
    assert r.status_code == 200, r.text
    return {s["login"]: s for s in r.json()["students"]}


def _export(client, headers, group, ids, **extra):
    return client.post("/web/accounts/rollout/export",
                       json={"group": group, "student_ids": list(ids), **extra},
                       headers=headers)


def _rows(resp):
    ws = load_workbook(io.BytesIO(resp.content)).active
    out = {}
    for row in ws.iter_rows(min_row=6, values_only=True):
        if row and isinstance(row[0], int):
            out[row[2]] = row[3]
    return out


def test_curator_sees_and_exports_only_own_groups(client):
    admin = make_admin(client)
    _student(client, admin, "a1", "К-1")
    _student(client, admin, "b1", "К-2")
    cur = _curator(client, admin, "cur1", ["К-1"])

    groups = client.get("/web/accounts/rollout/groups", headers=cur).json()["groups"]
    assert [g["name"] for g in groups] == ["К-1"], groups
    #Чужая группа — отказ на СЕРВЕРЕ, а не спрятанная кнопка.
    assert client.get("/web/accounts/rollout/students", params={"group": "К-2"},
                      headers=cur).status_code == 403
    assert _export(client, cur, "К-2", ["stud:b1"]).status_code == 403
    #Подставленный id студента из ЧУЖОЙ группы в запросе на СВОЮ — отбрасывается.
    r = _export(client, cur, "К-1", ["stud:b1"])
    assert r.status_code == 400, r.text
    #Обычный преподаватель (не куратор) и студент — не видят ничего.
    _student(client, admin, "st9", "К-1", password="studpass99")
    sh = {"Authorization": "Bearer " + client.post("/auth/login", json={
        "login": "st9", "password": "studpass99"}).json()["access_token"]}
    assert client.get("/web/accounts/rollout/groups", headers=sh).status_code == 403


def test_export_issues_passwords_that_work_and_repeat(client, data_key):
    admin = make_admin(client)
    _student(client, admin, "s1", "К-1")
    _student(client, admin, "s2", "К-1")
    ids = _ids(client, admin, "К-1")
    assert {v["state"] for v in ids.values()} == {"none"}

    r = _export(client, admin, "К-1", [v["id"] for v in ids.values()])
    assert r.status_code == 200, r.text
    first = _rows(r)
    assert set(first) == {"s1", "s2"}
    pw = first["s1"]
    assert pw and len(pw) >= 10
    for bad in "0Oo1lIi":
        assert bad not in pw, "в пароле не должно быть похожих символов"
    #Пароль из документа — рабочий.
    assert client.post("/auth/login", json={"login": "s1", "password": pw}).status_code == 200

    #Повторная выгрузка — те же пароли, а не новые.
    again = _rows(_export(client, admin, "К-1", [v["id"] for v in ids.values()]))
    assert again == first

    #В базе — шифротекст, а не пароль; в журнале аудита пароля нет.
    from app.db import SessionLocal
    from app.models import AuditEvent, UserIssuedCredential
    db = SessionLocal()
    try:
        row = db.get(UserIssuedCredential, "stud:s1")
        assert row.secret and pw not in row.secret and row.secret.startswith("gost1$")
        assert all(pw not in (e.detail or "") and pw not in (e.target or "")
                   for e in db.query(AuditEvent).all())
        assert db.query(AuditEvent).filter(
            AuditEvent.action == "credentials.rollout").count() == 2
    finally:
        db.close()


def test_export_does_not_kick_people_who_already_have_a_password(client):
    """Пароль задан до 4.0 — выгрузка его НЕ трогает без явного флага."""
    admin = make_admin(client)
    _student(client, admin, "old", "К-1", password="myoldpass1")
    ids = _ids(client, admin, "К-1")
    assert ids["old"]["state"] == "issued"   #набран администратором — это пароль колледжа

    #Имитация «задан до 4.0»: запись о выдаче отсутствует, хеш есть.
    from app.db import SessionLocal
    from app.models import UserIssuedCredential
    db = SessionLocal()
    try:
        db.delete(db.get(UserIssuedCredential, "stud:old"))
        db.commit()
    finally:
        db.close()
    assert _ids(client, admin, "К-1")["old"]["state"] == "preset"
    rows = _rows(_export(client, admin, "К-1", [ids["old"]["id"]]))
    assert "задан ранее" in rows["old"]
    assert client.post("/auth/login", json={"login": "old",
                                            "password": "myoldpass1"}).status_code == 200


def test_own_password_change_hides_the_issued_one(client, data_key):
    admin = make_admin(client)
    _student(client, admin, "s1", "К-1")
    r = client.post("/web/accounts/credential/reset", json={"login": "s1"}, headers=admin)
    assert r.status_code == 200, r.text
    pw = r.json()["password"]
    extra = client.get("/web/accounts/extra", params={"login": "s1"}, headers=admin).json()
    assert extra["credential"]["state"] == "issued" and extra["credential"]["password"] == pw

    tok = client.post("/auth/login", json={"login": "s1", "password": pw}).json()
    sh = {"Authorization": f"Bearer {tok['access_token']}", "X-Client": "web"}
    r = client.post("/me/password", json={"current": pw, "new": "свойПароль42"}, headers=sh)
    assert r.status_code == 200, r.text

    extra = client.get("/web/accounts/extra", params={"login": "s1"}, headers=admin).json()
    assert extra["credential"]["state"] == "changed"
    assert extra["credential"]["password"] == ""
    from app.db import SessionLocal
    from app.models import UserIssuedCredential, NotifyEvent
    db = SessionLocal()
    try:
        assert db.get(UserIssuedCredential, "stud:s1").secret == ""
        #Уведомление о смене есть — и пароля в нём нет.
        ev = db.query(NotifyEvent).filter(NotifyEvent.login == "s1",
                                          NotifyEvent.kind == "password_changed").all()
        assert ev and all("свойПароль42" not in (e.body or "") for e in ev)
    finally:
        db.close()
    #Выгрузка того, кто сменил пароль, его пароль не показывает и не сбрасывает.
    rows = _rows(_export(client, admin, "К-1", ["stud:s1"]))
    assert "сменил" in rows["s1"]
    assert client.post("/auth/login", json={"login": "s1",
                                            "password": "свойПароль42"}).status_code == 200


def test_the_guard_catches_a_path_it_does_not_know(client, data_key):
    """Сверка отпечатка хеша — не пожелание, а механизм: пароль сменён МИМО `forget`
    (как это сделал бы любой будущий путь смены), а стартовый всё равно не показывается.

    Обратный ход: без сверки в `issued_credentials.state` этот тест красный — старый
    пароль продолжал бы отдаваться администратору после смены."""
    admin = make_admin(client)
    _student(client, admin, "s1", "К-1")
    client.post("/web/accounts/credential/reset", json={"login": "s1"}, headers=admin)
    from app.db import SessionLocal
    from app.models import User, set_user_password
    db = SessionLocal()
    try:
        u = db.get(User, "stud:s1")
        set_user_password(u, "какой-то-другой-путь-1")
        db.commit()
    finally:
        db.close()
    extra = client.get("/web/accounts/extra", params={"login": "s1"}, headers=admin).json()
    assert extra["credential"] == {**extra["credential"], "state": "changed", "password": ""}


def test_curator_sees_credential_only_of_own_students(client):
    admin = make_admin(client)
    _student(client, admin, "mine", "К-1")
    _student(client, admin, "alien", "К-2")
    cur = _curator(client, admin, "cur2", ["К-1"])
    r = client.get("/web/accounts/extra", params={"login": "mine"}, headers=cur)
    assert r.status_code == 200, r.text
    assert "contacts" not in r.json(), "контакты студента куратору не отдаются"
    assert client.get("/web/accounts/extra", params={"login": "alien"},
                      headers=cur).status_code == 403
    assert client.post("/web/accounts/credential/reset", json={"login": "alien"},
                       headers=cur).status_code == 403
    #Контакты пишет только администратор.
    assert client.post("/web/accounts/extra", json={"login": "mine",
                                                    "admin_email": "x@mail.ru"},
                       headers=cur).status_code == 403
